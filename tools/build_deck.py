#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""สร้างสไลด์นำเสนอระบบแนะนำหนังสือ (16 สไลด์) ด้วย python-pptx

    python tools/build_deck.py     # interpreter ที่มี python-pptx + pywin32 คือ `python` (3.11)
"""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches as In, Pt
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
A = os.path.join(ROOT, "deck_assets")
OUT = os.path.join(ROOT, "slides"); os.makedirs(OUT, exist_ok=True)
PPTX = os.path.join(OUT, "นำเสนอระบบแนะนำหนังสือ_007.pptx")

SW, SH = 13.333, 7.5
BG, CARD, CARD2, BORDER = "#0E1626", "#172136", "#1E2A45", "#2B3A5C"
ORANGE, CYAN, GREEN, YELLOW, RED = "#FF8833", "#38BDF8", "#34D399", "#FBBF24", "#F87171"
TEXT, MUTED = "#E9EFFA", "#9FB0CB"
FONT = "Leelawadee UI"
GID = "247492673b9b3a74dd39bd7d442d398c"
COLAB = f"https://colab.research.google.com/gist/Nasak16/{GID}/BookRecommender_Neo4j_007.ipynb"


def hexc(h):
    return RGBColor.from_string(h.lstrip("#"))


prs = Presentation()
prs.slide_width, prs.slide_height = In(SW), In(SH)
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = hexc(BG)
    return s


def rect(s, x, y, w, h, fill=CARD, line=None, radius=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = s.shapes.add_shape(shape, In(x), In(y), In(w), In(h))
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = hexc(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = hexc(line); sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius if radius is not None else 0.06
    sh.text_frame.text = ""
    return sh


def tb(s, x, y, w, h, runs, size=14, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
       space_after=4, line_spacing=1.06, anchor=MSO_ANCHOR.TOP):
    """runs: str หรือ list ของ (ข้อความ, size, color, bold) หรือ list ของ list (ย่อหน้า)"""
    box = s.shapes.add_textbox(In(x), In(y), In(w), In(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) and runs and isinstance(runs[0], list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        items = para if isinstance(para, list) else [para]
        for it in items:
            if isinstance(it, tuple):
                text, sz, col, bd = (list(it) + [size, color, bold])[:4]
            else:
                text, sz, col, bd = it, size, color, bold
            r = p.add_run(); r.text = text
            r.font.size = Pt(sz); r.font.color.rgb = hexc(col)
            r.font.bold = bd; r.font.name = FONT
    return box


def header(s, kicker, title, num=None):
    rect(s, 0.62, 0.46, 0.075, 0.62, ORANGE, shape=MSO_SHAPE.RECTANGLE)
    tb(s, 0.82, 0.34, 10.8, 0.3, kicker, 11.5, ORANGE, True)
    tb(s, 0.82, 0.6, 11.4, 0.55, title, 27, TEXT, True)
    footer(s, num)


def footer(s, num=None):
    rect(s, 0, SH - 0.34, SW, 0.34, CARD2, shape=MSO_SHAPE.RECTANGLE)
    tb(s, 0.62, SH - 0.31, 9.5, 0.26, "ระบบแนะนำหนังสือ · Neo4j + Streamlit · จัดทำโดย รหัส 007",
       9, MUTED)
    if num:
        tb(s, SW - 1.15, SH - 0.31, 0.55, 0.26, str(num), 9.5, ORANGE, True, PP_ALIGN.RIGHT)


def pic_cover(s, path, x, y, w, h, border=True):
    """วางภาพแบบ cover-fit (ตัดขอบให้เต็มกรอบ ไม่ยืด)"""
    iw, ih = Image.open(path).size
    box_ar, img_ar = w / h, iw / ih
    pic = s.shapes.add_picture(path, In(x), In(y), In(w), In(h))
    if img_ar > box_ar:                      # ภาพกว้างเกิน → ตัดซ้าย/ขวา
        crop = (1 - box_ar / img_ar) / 2
        pic.crop_left = pic.crop_right = crop
    else:                                    # ภาพสูงเกิน → ตัดบน/ล่าง
        crop = (1 - img_ar / box_ar) / 2
        pic.crop_top = pic.crop_bottom = crop
    if border:
        pic.line.color.rgb = hexc(BORDER); pic.line.width = Pt(1)
    return pic


def pic_fit(s, path, x, y, w, h):
    """วางภาพแบบ contain (เห็นครบทั้งภาพ)"""
    iw, ih = Image.open(path).size
    ar = iw / ih
    if ar > w / h:
        ww, hh = w, w / ar
    else:
        hh, ww = h, h * ar
    return s.shapes.add_picture(path, In(x + (w - ww) / 2), In(y + (h - hh) / 2),
                                In(ww), In(hh))


def code_box(s, x, y, w, h, lines, size=11.5, title=None):
    rect(s, x, y, w, h, "#0B1220", BORDER, radius=0.04)
    yy = y + 0.12
    if title:
        tb(s, x + 0.2, yy, w - 0.4, 0.24, title, 10.5, CYAN, True)
        yy += 0.3
    tb(s, x + 0.2, yy, w - 0.4, h - (yy - y) - 0.15,
       [[(l, size, "#C9E4FF", False)] for l in lines], space_after=1, line_spacing=1.0)


def bullet_list(s, x, y, w, items, size=14, gap=0.06, dot=ORANGE):
    yy = y
    for it in items:
        rect(s, x, yy + 0.1, 0.085, 0.085, dot, shape=MSO_SHAPE.OVAL)
        box = tb(s, x + 0.26, yy, w - 0.26, 0.4, it, size, TEXT, space_after=2)
        lines = max(1, int((len(it) * size * 0.52) / (w * 96)) + 1)
        yy += 0.24 * lines + gap + 0.1
    return yy


def metric(s, x, y, w, h, value, label, color=ORANGE):
    rect(s, x, y, w, h, CARD, BORDER)
    tb(s, x, y + 0.16, w, 0.5, value, 26, color, True, PP_ALIGN.CENTER)
    tb(s, x, y + 0.72, w, 0.4, label, 10.5, MUTED, False, PP_ALIGN.CENTER)


# ---------------------------------------------------------------- 1 ปก
s = slide()
rect(s, 0.9, 1.05, 0.1, 1.5, ORANGE, shape=MSO_SHAPE.RECTANGLE)
tb(s, 1.25, 1.0, 11.0, 0.4, "งาน: พัฒนาระบบแนะนำเป็นระบบของตัวเอง", 14, CYAN, True)
tb(s, 1.25, 1.4, 11.2, 1.5, "ระบบแนะนำหนังสือ", 54, TEXT, True)
tb(s, 1.25, 2.55, 11.2, 0.6, "BOOK RECOMMENDER SYSTEM · Neo4j (Graph DB) + Streamlit", 20,
   ORANGE, True)
tb(s, 1.25, 3.35, 11.2, 1.0,
   [[("ข้อมูลของเราเอง 12 คน × 28 เล่ม × 51 ความชอบ", 15, TEXT, False)],
    [("ผลการแนะนำแสดง “ภาพปกหนังสือจริง” ทุกครั้ง", 15, TEXT, False)],
    [("มีทั้งโน๊ตบุ๊กอธิบายวิธีทำ (Colab), เว็บแอปสาธิต และสไลด์ชุดนี้บน GitHub", 15, TEXT, False)]],
   space_after=4)
for i, (v, l, c) in enumerate([("12", "ผู้ใช้ในระบบ", CYAN), ("28", "หนังสือ (14 หมวด)", YELLOW),
                               ("51", "ความชอบ", GREEN), ("3+1", "วิธีให้คะแนน", ORANGE)]):
    metric(s, 1.25 + i * 2.6, 5.15, 2.3, 1.25, v, l, c)
tb(s, 1.25, 6.7, 11.2, 0.4, "จัดทำโดย: รหัส 007 · วิชา ______________ · วันที่ ____/____/______",
   12, MUTED)

# ---------------------------------------------------------------- 2 สารบัญ
s = slide(); header(s, "AGENDA", "หัวข้อที่จะนำเสนอ", 2)
agenda = [("01", "โจทย์และสิ่งที่ส่ง", "งาน 4 ข้อของอาจารย์ → ส่งครบอะไรบ้าง"),
          ("02", "แนวคิดระบบแนะนำบนกราฟ", "ทำไมข้อมูล “ใครชอบอะไร” ต้องใช้กราฟ + การเดิน 3 hop"),
          ("03", "ข้อมูลและสถาปัตยกรรม", "ข้อมูลชุดของเรา + Neo4j + Streamlit + GitHub"),
          ("04", "ระบบทำงานจริง", "ผลการแนะนำ 3 วิธี + ภาพปก + เหตุผลย้อนหลังได้"),
          ("05", "สาธิตการใช้งาน (เดโมสด)", "เพิ่มความชอบแล้วคำแนะนำเปลี่ยนทันทีโดยไม่ต้องเทรนใหม่"),
          ("06", "การทดสอบและลิงก์งาน", "ผลตรวจสอบจริง + Colab + GitHub + หน้า index")]
for i, (n, t, d) in enumerate(agenda):
    y = 1.5 + i * 0.92
    rect(s, 0.82, y, 11.7, 0.78, CARD, BORDER)
    tb(s, 1.0, y + 0.14, 0.8, 0.4, n, 22, ORANGE, True)
    tb(s, 1.85, y + 0.09, 4.6, 0.35, t, 15.5, TEXT, True)
    tb(s, 1.85, y + 0.42, 10.4, 0.3, d, 11.5, MUTED)

# ---------------------------------------------------------------- 3 โจทย์
s = slide(); header(s, "โจทย์ & สิ่งที่ส่ง", "งานที่อาจารย์สั่ง 4 ข้อ — ส่งครบทุกข้อ", 3)
items = [("1", "Present ด้วย PowerPoint + สาธิตการใช้งานระบบ",
          "สไลด์ 16 หน้า (ไฟล์นี้) + เดโมสดในแอป + ใส่ภาพในระบบแนะนำจริง (หน้าปกหนังสือ)"),
         ("2", "ข้อมูล PowerPoint ไว้ใน GitHub", "ใส่ทั้ง .pptx และ .pdf ในโฟลเดอร์ slides/ ของ repo"),
         ("3", "เอาการบ้านทุกอันไว้ใน GitHub", "ทุกชิ้นงานมี repo ของตัวเอง + โฟลเดอร์กลาง homework รวมลิงก์"),
         ("4", "จัดหน้า index ให้ลิงก์กับงานทุกอัน", "หน้าเว็บ index.html บน GitHub Pages ลิงก์ทุก repo")]
for i, (n, t, d) in enumerate(items):
    y = 1.5 + i * 1.28
    rect(s, 0.82, y, 11.7, 1.1, CARD, BORDER)
    rect(s, 0.82, y, 0.075, 1.1, ORANGE, shape=MSO_SHAPE.RECTANGLE)
    tb(s, 1.05, y + 0.18, 0.5, 0.5, n, 24, ORANGE, True)
    tb(s, 1.75, y + 0.14, 10.5, 0.4, t, 16, TEXT, True)
    tb(s, 1.75, y + 0.58, 10.5, 0.4, d, 12, MUTED)

# ---------------------------------------------------------------- 4 แนวคิด
s = slide(); header(s, "แนวคิด", "ทำไมต้องใช้ฐานข้อมูลกราฟ?", 4)
tb(s, 0.82, 1.35, 6.1, 0.6,
   "ข้อมูล “ใครชอบหนังสือเล่มไหน” เป็นเรื่องของความสัมพันธ์ — ถ้าเก็บเป็นตาราง "
   "จะต้อง JOIN หลายชั้นกว่าจะได้คำตอบ แต่ในกราฟเขียนครั้งเดียวจบ", 13, TEXT)
code_box(s, 0.82, 2.15, 6.1, 2.35,
         ["MATCH (me:User {user: 'สมชาย'})-[:LIKES]->",
          "      (shared:Book)<-[:LIKES]-(other:User)",
          "MATCH (other)-[:LIKES]->(rec:Book)",
          "WHERE NOT (me)-[:LIKES]->(rec)",
          "RETURN rec.title, count(DISTINCT other) AS votes",
          "ORDER BY votes DESC"],
         title="คำแนะนำทั้งระบบเขียนด้วย Cypher ไม่กี่บรรทัด")
tb(s, 0.82, 4.6, 6.1, 0.4, "ขั้นตอนการเดิน 3 hop", 13, ORANGE, True)
tb(s, 0.82, 4.95, 6.1, 1.7,
   [[("1. ดูหนังสือที่ผู้ใช้ชอบ", 12, TEXT, False)],
    [("2. หาคนอื่นที่ชอบเล่มเดียวกัน = รสนิยมใกล้กัน", 12, TEXT, False)],
    [("3. ดูเล่มอื่นที่คนกลุ่มนั้นชอบ (ที่เรายังไม่ชอบ)", 12, TEXT, False)],
    [("4. นับ/ให้คะแนน → เรียงอันดับ = คำแนะนำ", 12, TEXT, False)]], space_after=3)
rect(s, 7.2, 1.35, 5.3, 5.25, CARD, BORDER)
tb(s, 7.5, 1.5, 4.8, 0.4, "ตัวอย่างจริง: สมชาย", 16, ORANGE, True)
tb(s, 7.5, 1.95, 4.8, 4.4,
   [[("ชอบอยู่แล้ว 4 เล่ม", 12.5, MUTED, True)],
    [("Harry Potter · The Hobbit · LOTR · Da Vinci Code", 12, TEXT, False)],
    [("คนรสนิยมใกล้ (Jaccard)", 12.5, MUTED, True)],
    [("นรินทร์ 0.500 (ชอบร่วม 3 เล่ม)", 12, TEXT, False)],
    [("มะลิ 0.333 (ชอบร่วม 2 เล่ม)", 12, TEXT, False)],
    [("พลอย 0.143 (ชอบร่วม 1 เล่ม)", 12, TEXT, False)],
    [("ระบบแนะนำ", 12.5, MUTED, True)],
    [("The Hunger Games — 0.833", 13, GREEN, True)],
    [("Dune — 0.500", 12, TEXT, False)],
    [("The Lightning Thief — 0.333", 12, TEXT, False)],
    [("เพราะมะลิกับนรินทร์ต่างชอบ Hunger Games ทั้งคู่", 11.5, MUTED, False)]], space_after=4)

# ---------------------------------------------------------------- 5 ข้อมูล
s = slide(); header(s, "ข้อมูลชุดของเรา", "12 คน × 28 เล่ม × 51 ความชอบ (เก็บเอง ไม่ใช้ dataset สำเร็จรูป)", 5)
for i, (v, l, c) in enumerate([("12", "ผู้ใช้ (เพื่อนในห้อง)", CYAN), ("28", "หนังสือ", YELLOW),
                               ("14", "หมวดหนังสือ", GREEN), ("51", "ความสัมพันธ์ LIKES", ORANGE)]):
    metric(s, 0.82 + i * 1.62, 1.35, 1.45, 1.2, v, l, c)
tb(s, 0.82, 2.75, 6.0, 1.5,
   [[("ปกหนังสือทุกเล่มดึงจาก Open Library (ฟรี ไม่ต้องใช้ API key) "
      "แล้วเก็บไฟล์ไว้ในโปรเจกต์", 12.5, TEXT, False)],
    [("→ ภาพยังแสดงได้แม้ไม่มีอินเทอร์เน็ต และเปิดใน Colab ก็เห็นภาพ", 12.5, MUTED, False)],
    [("หมวด: แฟนตาซี · ไซไฟ · สืบสวน · พัฒนาตัวเอง · การเงิน · การ์ตูน · คลาสสิก ฯลฯ",
      11.5, MUTED, False)]], space_after=5)
tb(s, 0.82, 4.5, 6.0, 0.4, "ขั้นตอนเตรียมข้อมูล", 13, ORANGE, True)
bullet_list(s, 0.82, 4.9, 6.0, [
    "ค้นหาแต่ละเล่มใน Open Library API → ได้ ISBN + รหัสภาพปก",
    "ดาวน์โหลดภาพปกเก็บเป็น asset ของ repo (ตรวจด้วยตาว่าตรงเรื่องทุกเล่ม)",
    "กำหนดความชอบของ 12 คนให้มีกลุ่มรสนิยมชัดเจน และไม่มีหนังสือเล่มไหนไร้เจ้าของ",
], size=12, gap=0.02)
pic_cover(s, os.path.join(A, "cover_sheet.png"), 7.05, 1.35, 5.5, 5.4)

# ---------------------------------------------------------------- 6 สถาปัตยกรรม
s = slide(); header(s, "สถาปัตยกรรม", "ระบบประกอบด้วย 4 ส่วน และทุกส่วนอยู่บน GitHub", 6)
arch = [("Colab Notebook", "ทดลอง + อธิบายวิธีทำ\nCypher ทีละขั้น มีผลรันจริงฝังในไฟล์", CYAN, 0.5),
        ("Neo4j", "ฐานข้อมูลกราฟ\n12 User · 28 Book · 14 Genre · 51 LIKES", ORANGE, 3.62),
        ("Streamlit App", "หน้าจอสาธิต 5 แท็บ\nการ์ดภาพปก · กราฟ · เดโมสด", GREEN, 6.74),
        ("GitHub + Pages", "repo ระบบ + โฟลเดอร์ homework\nหน้า index รวมลิงก์งานทุกชิ้น", YELLOW, 9.86)]
for name, desc, col, x in arch:
    rect(s, x, 1.5, 2.85, 2.5, CARD, BORDER)
    rect(s, x, 1.5, 2.85, 0.09, col, shape=MSO_SHAPE.RECTANGLE)
    tb(s, x + 0.22, 1.75, 2.45, 0.5, name, 15.5, col, True)
    tb(s, x + 0.22, 2.35, 2.45, 1.5, [[(l, 11.5, TEXT, False)] for l in desc.split("\n")],
       space_after=3)
for x in (3.37, 6.49, 9.61):
    tb(s, x, 2.55, 0.24, 0.4, "▶", 16, MUTED, True, PP_ALIGN.CENTER)
rect(s, 0.82, 4.35, 11.7, 2.3, CARD, BORDER)
tb(s, 1.05, 4.5, 11.2, 0.4, "การไหลของข้อมูล", 14, ORANGE, True)
tb(s, 1.05, 4.95, 11.2, 1.6,
   [[("① เตรียมข้อมูล", 12.5, CYAN, True), (" — สคริปต์ดึงข้อมูล + ภาพปกจาก Open Library → data/books.json", 12, TEXT, False)],
    [("② โหลดเข้า Neo4j", 12.5, ORANGE, True), (" — MERGE โหนด/ความสัมพันธ์ + สร้าง constraint กันข้อมูลซ้ำ", 12, TEXT, False)],
    [("③ ให้บริการคำแนะนำ", 12.5, GREEN, True), (" — Cypher 3 วิธี + วิธีผสม (ถ่วงน้ำหนัก Jaccard)", 12, TEXT, False)],
    [("④ แสดงผล", 12.5, YELLOW, True), (" — Streamlit แสดงการ์ดภาพปก + กราฟ + เดโมสดแก้ข้อมูลในฐานข้อมูลจริง", 12, TEXT, False)]],
   space_after=5)

# ---------------------------------------------------------------- 7 สคีมา
s = slide(); header(s, "โครงสร้างข้อมูล", "สคีมากราฟและคำสั่งที่ใช้จริง", 7)
rect(s, 0.82, 1.35, 5.6, 1.5, CARD, BORDER)
tb(s, 1.05, 1.5, 5.2, 1.2,
   [[("(User)", 16, ORANGE, True), ("  -[:LIKES]->  ", 13, MUTED, False), ("(Book)", 16, YELLOW, True),
     ("  -[:IN_GENRE]->  ", 13, MUTED, False), ("(Genre)", 16, GREEN, True)],
    [("บุคลิกของข้อมูล: ผู้ใช้ → ความชอบ → หนังสือ → หมวด", 11.5, MUTED, False)]], space_after=6)
tb(s, 0.82, 3.0, 5.6, 0.35, "ข้อดีที่ได้จากโครงสร้างนี้", 13, ORANGE, True)
bullet_list(s, 0.82, 3.4, 5.6, [
    "ถามจากโหนดไหนก็ได้ ไม่ต้อง JOIN หลายตาราง",
    "เพิ่มคน/หนังสือใหม่ = เพิ่มโหนด ไม่ต้องแก้โครงสร้าง",
    "โหนด Genre เป็นทั้ง “สัญญาณที่สอง” และตัวช่วยแก้ cold start",
], size=12)
code_box(s, 6.75, 1.35, 5.8, 2.55,
         ["MATCH (u:User)-[:LIKES]->(b:Book)",
          "RETURN u.user AS user, count(b) AS books_liked",
          "ORDER BY books_liked DESC",
          "",
          "// จำนวนจริงในฐานข้อมูล",
          "// User 12 · Book 28 · Genre 14 · LIKES 51"],
         title="Cypher ตัวอย่าง")
code_box(s, 6.75, 4.1, 5.8, 2.55,
         ["CREATE CONSTRAINT user_name IF NOT EXISTS",
          "FOR (u:User) REQUIRE u.user IS UNIQUE;",
          "",
          "UNWIND $rows AS row",
          "MERGE (u:User {user: row.user})",
          "SET u.age = row.age, u.group = row.group;"],
         title="กันข้อมูลซ้ำตอนโหลด (constraint + MERGE)")

# ---------------------------------------------------------------- 8 ผลการแนะนำ
s = slide(); header(s, "ระบบทำงานจริง (1)", "หน้าแนะนำหนังสือ — มีภาพปกในระบบแนะนำตามโจทย์", 8)
pic_cover(s, os.path.join(A, "app_reco.png"), 0.82, 1.32, 8.15, 5.4)
cards = [("🎯", "การ์ดแนะนำ", "ภาพปก + ชื่อเรื่อง + ผู้เขียน + หมวด", GREEN),
         ("⭐", "คะแนน", "มาจากความคล้ายของคนที่ชอบเล่มนี้", YELLOW),
         ("🧠", "เหตุผล", "บอกชื่อเพื่อนและเล่มที่ชอบร่วมกัน", CYAN),
         ("📚", "ข้อมูลตั้งต้น", "แสดงหนังสือที่ผู้ใช้ชอบอยู่แล้ว", ORANGE)]
for i, (ic, t, d, c) in enumerate(cards):
    y = 1.32 + i * 1.38
    rect(s, 9.2, y, 3.3, 1.25, CARD, BORDER)
    tb(s, 9.4, y + 0.12, 0.5, 0.4, ic, 18, c, True)
    tb(s, 9.95, y + 0.14, 2.4, 0.35, t, 13.5, c, True)
    tb(s, 9.95, y + 0.52, 2.45, 0.6, d, 11, MUTED)

# ---------------------------------------------------------------- 9 วิธีให้คะแนน
s = slide(); header(s, "ระบบทำงานจริง (2)", "ให้คะแนน 3 วิธี + วิธีผสม — ผลจริงของผู้ใช้ “สมชาย”", 9)
cols = [("นับโหวต", "1 คน = 1 เสียง", ["The Hunger Games 2 โหวต", "Dune 1 โหวต", "อื่น ๆ 1 โหวต"],
         CYAN, "ปัญหาที่เจอ: เพื่อนทุกคนมีน้ำหนักเท่ากัน แม้รสนิยมจะใกล้ไม่เท่ากัน"),
        ("ถ่วงน้ำหนัก (Jaccard)", "คะแนน = Σ ความคล้าย", ["The Hunger Games 0.833", "Dune 0.500",
                                                          "The Lightning Thief 0.333"],
         GREEN, "แก้ปัญหาเดิม: คนรสนิยมใกล้สุดมีน้ำหนักมากสุด"),
        ("ตามหมวด + ผสม", "เพื่อน 60% + หมวด 40%", ["The Hunger Games 0.600",
                                                     "And Then There Were None 0.503",
                                                     "Dune 0.400"],
         ORANGE, "ช่วยผู้ใช้ใหม่ที่ยังไม่มีความคล้ายกับใคร (cold start)")]
for i, (t, sub, rows, c, note) in enumerate(cols):
    x = 0.82 + i * 3.95
    rect(s, x, 1.4, 3.65, 4.15, CARD, BORDER)
    rect(s, x, 1.4, 3.65, 0.09, c, shape=MSO_SHAPE.RECTANGLE)
    tb(s, x + 0.2, 1.6, 3.3, 0.4, t, 15.5, c, True)
    tb(s, x + 0.2, 2.0, 3.3, 0.3, sub, 11.5, MUTED)
    for j, r in enumerate(rows):
        rect(s, x + 0.2, 2.4 + j * 0.52, 3.25, 0.44, CARD2, BORDER)
        tb(s, x + 0.34, 2.48 + j * 0.52, 3.0, 0.3, r, 11.5, TEXT)
    tb(s, x + 0.2, 4.2, 3.25, 1.2, note, 11, MUTED)
tb(s, 0.82, 5.75, 11.7, 0.9,
   [[("สูตรถ่วงน้ำหนัก:  ", 12.5, ORANGE, True),
     ("Jaccard(เรา, เขา) = |หนังสือที่ชอบร่วมกัน| ÷ |หนังสือของสองคนรวมกัน|  "
      "→ ใช้เป็นน้ำหนักโหวตของแต่ละคน แล้วบวกกันเป็นคะแนนของเล่มนั้น", 12, TEXT, False)],
    [("ทุกคำแนะนำอธิบายย้อนหลังได้: บอกได้ว่าเล่มนี้มาจากเพื่อนคนไหน และเพราะคุณชอบเล่มอะไร", 12, MUTED, False)]])

# ---------------------------------------------------------------- 10 เดโม
s = slide(); header(s, "สาธิตการใช้งาน", "เดโมสด: เพิ่มความชอบ → คำแนะนำเปลี่ยนทันที (ไม่ต้องเทรนโมเดลใหม่)", 10)
pic_cover(s, os.path.join(A, "app_demo.png"), 0.82, 1.35, 7.6, 4.6)
rect(s, 8.65, 1.35, 3.85, 4.6, CARD, BORDER)
tb(s, 8.88, 1.5, 3.4, 0.4, "สาธิตตามลำดับนี้", 14, ORANGE, True)
steps = ["เปิดหน้า “เดโมสด: เพิ่มความชอบ”",
         "เลือกเล่มที่ต้องการเพิ่ม (เช่น The Martian)",
         "กด “➕ เพิ่มความชอบ” → เขียนลง Neo4j ทันที",
         "ตารางเทียบ ก่อน/หลัง จะโชว์ว่าคำแนะนำเปลี่ยนเป็นเล่มใหม่",
         "ลบข้อมูลที่เพิ่ม เพื่อคืนสภาพเดิม (กดปุ่มลบได้)",
         "สลับวิธีให้คะแนนเพื่อโชว์ผลต่างของแต่ละวิธี"]
yy = 1.95
for i, st in enumerate(steps):
    tb(s, 8.88, yy, 0.35, 0.3, str(i + 1) + ".", 12, CYAN, True)
    box = tb(s, 9.25, yy, 3.05, 0.75, st, 11.5, TEXT)
    yy += 0.62
tb(s, 0.82, 6.15, 11.7, 0.5,
   [[("จุดที่ทำให้ระบบนี้ “เป็นระบบของตัวเอง”: อัลกอริทึมของเรา + ข้อมูลของเรา + "
      "หน้าจอของเรา และเขียนคำสั่งกราฟเองทั้งหมด", 12, MUTED, False)]])

# ---------------------------------------------------------------- 11 คลังหนังสือ
s = slide(); header(s, "ระบบทำงานจริง (3)", "คลังหนังสือ 28 เล่ม — ค้นหาได้ และแสดงปกครบทุกเล่ม", 11)
pic_cover(s, os.path.join(A, "app_catalog.png"), 0.82, 1.32, 7.9, 5.4)
tb(s, 9.0, 1.45, 3.5, 0.4, "สิ่งที่หน้าแสดง", 14, ORANGE, True)
bullet_list(s, 9.0, 1.9, 3.5, [
    "ค้นหาด้วยชื่อเรื่อง / นักเขียน / หมวด",
    "แต่ละการ์ด: ปก + ชื่อ + หมวด + จำนวนคนที่ชอบ",
    "ใช้ตรวจว่าข้อมูลในกราฟครบ 28 เล่มจริง",
], size=11.5, gap=0.0)
rect(s, 9.0, 3.6, 3.5, 3.1, CARD, BORDER)
tb(s, 9.2, 3.75, 3.1, 0.4, "เทคโนโลยีที่ใช้", 13, CYAN, True)
tb(s, 9.2, 4.2, 3.1, 2.35,
   [[("Neo4j", 12.5, TEXT, True), (" — ฐานข้อมูลกราฟ (Cypher)", 11.5, MUTED, False)],
    [("Streamlit", 12.5, TEXT, True), (" — ส่วนติดต่อผู้ใช้บนเว็บ (ฟรี)", 11.5, MUTED, False)],
    [("networkx + matplotlib", 12.5, TEXT, True), (" — วาดกราฟความสัมพันธ์", 11.5, MUTED, False)],
    [("Open Library API", 12.5, TEXT, True), (" — ข้อมูลและภาพปก (ฟรี)", 11.5, MUTED, False)],
    [("Colab + gist", 12.5, TEXT, True), (" — โน๊ตบุ๊กอธิบายวิธีทำ", 11.5, MUTED, False)]],
   space_after=6)

# ---------------------------------------------------------------- 12 กราฟ
s = slide(); header(s, "ระบบทำงานจริง (4)", "ภาพกราฟความสัมพันธ์ที่ระบบวาดเอง (ผู้ใช้ ↔ หนังสือ)", 12)
pic_fit(s, os.path.join(A, "graph_somchai.png"), 0.7, 1.35, 11.9, 4.5)
legend = [("ส้ม", "ผู้ใช้ที่เลือก (สมชาย)", ORANGE), ("ฟ้า", "ผู้ใช้รสนิยมใกล้", CYAN),
          ("เหลือง", "หนังสือที่ชอบแล้ว", YELLOW), ("เขียว", "หนังสือที่แนะนำ", GREEN),
          ("เส้นประ", "เส้นทาง 3 hop", GREEN)]
for i, (c, t, col) in enumerate(legend):
    x = 0.82 + i * 2.42
    rect(s, x, 6.0, 2.25, 0.62, CARD, BORDER)
    rect(s, x + 0.16, 6.22, 0.16, 0.16, col, shape=MSO_SHAPE.OVAL)
    tb(s, x + 0.42, 6.14, 1.75, 0.4, c + " = " + t, 10, MUTED)
tb(s, 0.82, 5.95 - 0.0, 11.7, 0.4, "", 10, MUTED)   # เว้นวรรคกันชนกับรูป

# ---------------------------------------------------------------- 13 cold start
s = slide(); header(s, "ทดสอบระบบ", "Cold start และการเทียบกับ baseline ยอดนิยม", 13)
rect(s, 0.82, 1.4, 5.7, 4.4, CARD, BORDER)
tb(s, 1.05, 1.55, 5.2, 0.4, "ผู้ใช้ใหม่ที่ชอบเล่มเดียว (ที่ยังไม่มีใครชอบ)", 14, ORANGE, True)
tb(s, 1.05, 2.05, 5.25, 3.5,
   [[("ผู้ใช้ใหม่ชอบ: One-Punch Man, Vol. 1 (หมวดการ์ตูน)", 12, TEXT, True)],
    [("ผลจากสัญญาณเพื่อน", 12, MUTED, True), (" → ว่างเปล่า (cold start จริง)", 12, RED, False)],
    [("ผลจากสัญญาณหมวดหนังสือ", 12, MUTED, True),
     (" → Death Note Vol. 1 · One Piece Vol. 1", 12, GREEN, True)],
    [("บทสรุป: ถ้าใช้เฉพาะ collaborative filtering ผู้ใช้แบบนี้จะไม่ได้คำแนะนำเลย "
      "แต่เพราะมีโหนด Genre ระบบจึงยังแนะนำได้ — และเมื่อผู้ใช้ใหม่มีความชอบเพิ่ม "
      "คำแนะนำจะเปลี่ยนไปตามเพื่อนที่รสนิยมตรงกัน", 12, TEXT, False)]], space_after=6)
rect(s, 6.9, 1.4, 5.62, 4.4, CARD, BORDER)
tb(s, 7.15, 1.55, 5.1, 0.4, "เทียบกับ baseline ‘ยอดนิยม’", 14, ORANGE, True)
code_box(s, 7.15, 2.05, 5.1, 1.5,
         ["MATCH (b:Book)<-[:LIKES]-(u:User)",
          "RETURN b.title, count(u) AS likes",
          "ORDER BY likes DESC LIMIT 5"],
         title="แนะนำแบบไม่ดูว่าใครเป็นใคร")
tb(s, 7.15, 3.7, 5.1, 2.0,
   [[("Harry Potter · The Hobbit · The Hunger Games (3 คน)", 11.5, TEXT, False)],
    [("ทุกคนจะได้ลิสต์เดียวกันเสมอ", 12, RED, True)],
    [("ระบบของเราดีกว่า เพราะดูความคล้ายของแต่ละคน", 12, GREEN, True),
     (" — สมชายได้ Hunger Games แต่ ‘ธนกร’ ได้ Dune/Frankenstein/The Martian "
      "และ ‘มินต์’ ได้สายพัฒนาตัวเอง", 11.5, MUTED, False)]], space_after=5)
tb(s, 0.82, 5.95, 11.7, 0.6,
   [[("ผลรันจริงจากโน๊ตบุ๊ก (ฝัง output ในไฟล์ส่ง): ทุกเซลล์รันผ่าน · "
      "การ์ดภาพปกปรากฏ 16 จุด · ฐานข้อมูลกลับสู่สภาพ 12/28/51 หลังทดลอง cold start", 11.5, MUTED, False)]])

# ---------------------------------------------------------------- 14 การทดสอบ
s = slide(); header(s, "การทดสอบ", "ผลตรวจสอบจริง (ไม่ใช่คำอธิบาย)", 14)
tests = [("โหลดข้อมูลเข้า Neo4j", "tools/load_neo4j.py",
          "12 ผู้ใช้ · 28 หนังสือ · 51 LIKES · 14 Genre", GREEN),
         ("ความสอดคล้อง 2 backend", "tools/consistency_test.py",
          "Cypher (Neo4j) กับ backend สำรองให้ผลตรงกัน 12 คน × 5 การตรวจ", GREEN),
         ("โน๊ตบุ๊กทั้งเล่ม", "py -3.13 tools/verify_ipynb_cells.py",
          "VERDICT: PASS — ทุก code cell รันได้ (12 เซลล์)", GREEN),
         ("ภาพปกครบถ้วน", "contact sheet + ตรวจด้วยสายตา",
          "28/28 เล่มมีปกจริง (แก้ 2 เล่มที่เป็นภาพเปล่า)", YELLOW),
         ("หน้าจอแอปจริง", "tools/shots.py (Playwright)",
          "ถ่าย 6 ภาพ ทุกแท็บแสดงปกครบ ไม่มีภาพพัง", GREEN)]
for i, (t, cmd, res, c) in enumerate(tests):
    y = 1.4 + i * 1.0
    rect(s, 0.82, y, 11.7, 0.88, CARD, BORDER)
    rect(s, 0.82, y, 0.075, 0.88, c, shape=MSO_SHAPE.RECTANGLE)
    tb(s, 1.05, y + 0.12, 3.3, 0.35, t, 13.5, TEXT, True)
    tb(s, 1.05, y + 0.48, 4.3, 0.3, cmd, 10, CYAN)
    tb(s, 5.5, y + 0.2, 6.8, 0.5, res, 11.5, c)
tb(s, 0.82, 6.45, 11.7, 0.4,
   "หมายเหตุ: ผลการรันในโน๊ตบุ๊กมาจาก Neo4j ที่ทดสอบในเครื่อง (เซลล์เชื่อมต่อยังคงเป็นบัญชีของเรา "
   "พร้อมรหัสผ่านที่กรอกเองตอนรันบน Colab)", 10.5, MUTED)

# ---------------------------------------------------------------- 15 ลิงก์/ใช้งาน
s = slide(); header(s, "การใช้งานและลิงก์", "เปิดดูได้ทันที / รันเองได้ใน 3 คำสั่ง", 15)
boxes = [("โน๊ตบุ๊กอธิบายวิธีทำ (Colab)", COLAB, CYAN),
         ("โค้ดระบบทั้งหมด (GitHub)", "https://github.com/Nasak16/book-recommender", GREEN),
         ("หน้า index รวมงานทุกชิ้น", "https://nasak16.github.io/homework/", YELLOW)]
for i, (t, url, c) in enumerate(boxes):
    y = 1.45 + i * 1.15
    rect(s, 0.82, y, 11.7, 1.0, CARD, BORDER)
    rect(s, 0.82, y, 0.075, 1.0, c, shape=MSO_SHAPE.RECTANGLE)
    tb(s, 1.05, y + 0.14, 4.6, 0.35, t, 13.5, TEXT, True)
    tb(s, 1.05, y + 0.52, 11.2, 0.35, url, 11, c)
code_box(s, 0.82, 5.1, 5.6, 1.7,
         ["py -3.13 -m pip install -r requirements.txt",
          "py -3.13 tools/load_neo4j.py <neo4j-uri> neo4j <password>",
          "py -3.13 -m streamlit run app.py"],
         title="รันในเครื่อง (3 คำสั่ง)")
tb(s, 6.75, 5.25, 5.8, 1.5,
   [[("ถ้าต่อ Neo4j ไม่ได้ แอปจะสลับไปใช้ backend สำรองอัตโนมัติ และแจ้งเตือนบนแถบข้าง "
      "— หน้าจอและภาพปกยังทำงานครบถ้วน", 11.5, MUTED, False)],
    [("รหัสผ่าน Neo4j ไม่ถูกอัปโหลด (อยู่ใน .gitignore) — "
      "ในโปรเจกต์มีแค่ไฟล์ตัวอย่าง secrets.toml.example", 11.5, MUTED, False)]], space_after=6)

# ---------------------------------------------------------------- 16 สรุป
s = slide(); header(s, "สรุป", "ระบบแนะนำของเรา + สิ่งที่ต่อยอดได้", 16)
rect(s, 0.82, 1.4, 5.7, 4.5, CARD, BORDER)
tb(s, 1.05, 1.55, 5.2, 0.4, "ได้อะไรจากงานนี้", 14, GREEN, True)
bullet_list(s, 1.05, 2.0, 5.2, [
    "ระบบแนะนำที่ใช้ได้จริงบนฐานข้อมูลกราฟ + ข้อมูลของเราเอง",
    "แสดงภาพปกจริงในผลการแนะนำ (ตามโจทย์ข้อ 1)",
    "คำแนะนำอธิบายย้อนหลังได้ และทดสอบ cold start แล้ว",
    "มีทั้งโน๊ตบุ๊ก เว็บแอปสไลด์ และหน้า index รวมงานบน GitHub",
], size=12, gap=0.02, dot=GREEN)
rect(s, 6.9, 1.4, 5.62, 4.5, CARD, BORDER)
tb(s, 7.15, 1.55, 5.1, 0.4, "ข้อจำกัด และก้าวต่อไป", 14, YELLOW, True)
bullet_list(s, 7.15, 2.0, 5.1, [
    "ข้อมูล 51 ความชอบยังน้อย (sparsity) — เก็บจากผู้ใช้จริงเพิ่มจะแม่นขึ้น",
    "ยังไม่มีคะแนนแบบ 1-5 ดาว จึงถ่วงน้ำหนักได้แค่ความคล้ายรสนิยม",
    "ต่อยอด: ใช้ tag ละเอียดกว่า 14 หมวด, เพิ่ม PageRank/community detection, "
    "ทำระบบสมัครสมาชิกจริงบน Streamlit Cloud",
], size=12, gap=0.02, dot=YELLOW)
tb(s, 0.82, 6.15, 11.7, 0.5, "ขอบคุณครับ — ยินดีตอบคำถาม / สาธิตสดได้ทันที", 14, ORANGE, True,
   PP_ALIGN.CENTER)

prs.save(PPTX)
print("บันทึกสไลด์:", PPTX, "| จำนวนสไลด์:", len(prs.slides))
print("ขนาดไฟล์:", os.path.getsize(PPTX) // 1024, "KB")
